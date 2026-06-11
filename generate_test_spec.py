#!/usr/bin/env python3
"""SoFit 단위테스트 명세서 생성기"""

import os, re
from pathlib import Path
from collections import defaultdict, Counter
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR   = "/Users/ohsaekbit/sboh/SoFit-backend"
OUTPUT     = os.path.join(BASE_DIR, "SoFit_단위테스트_명세서.xlsx")
ADMIN_TEST = os.path.join(BASE_DIR, "sofit-admin/src/test")
USER_TEST  = os.path.join(BASE_DIR, "sofit-user/src/test")

# ── ID 컨벤션 ────────────────────────────────────────────────────

LOAN_SUBDOMAIN = [
    # 긴 prefix 우선
    ("LoanApplicationReview",    "RVW"),
    ("LoanApplicationGrade",     "GRD"),
    ("LoanApplicationInfo",      "INF"),
    ("LoanApplicationService",   "APP"),
    ("LoanApplicationConverter", "APP"),
    ("LoanApplicationController","APP"),
    ("LoanApplication",          "APP"),
    ("LoanDecisionProcessor",    "DEC"),
    ("LoanDecisionTasklet",      "DEC"),
    ("LoanDecision",             "DEC"),
    ("LoanDashboard",            "DSB"),
    ("LoanExecution",            "EXC"),
    ("LoanProduct",              "PRD"),
    ("LoanStep",                 "STP"),
    ("LoanStatistics",           "STAT"),
    ("ManagerApproval",          "MGR"),
    ("MyBizDataDetail",          "MBZD"),
    # 나머지 → ETC
]

LAYER_MAP = [
    ("ControllerTest",  "C"),
    ("ServiceImplTest", "S"),
    ("ConverterTest",   "CV"),
    ("ClientTest",      "CL"),
    ("SchedulerTest",   "SCH"),
    ("FilterTest",      "FT"),
    ("ListenerTest",    "EV"),
    ("ProcessorTest",   "BT"),
    ("TaskletTest",     "BT"),
    ("UtilTest",        "UTL"),
    ("ServiceTest",     "S"),
    ("Test",            "UTL"),
]

PKG_DOMAIN = {
    "auth": "AUTH", "dev": "DEV",
    "mybiz": "MBZ", "notification": "NOTIF",
    "report": "RPT", "terms": "TERM", "user": "USR",
}

# ── 스타일 상수 ─────────────────────────────────────────────────

LAYER_COLOR = {
    "C":   "D6E4F0",
    "S":   "E8F5E9",
    "CV":  "FFF3E0",
    "CL":  "F3E5F5",
    "BT":  "FCE4EC",
    "SCH": "E0F7FA",
    "FT":  "FFF9C4",
    "EV":  "F1F8E9",
    "UTL": "F5F5F5",
}
HEADER_BG   = "1F4E79"
PASS_BG     = "C8E6C9"
PASS_FG     = "2E7D32"
TOTAL_BG    = "E3E3E3"

DOMAIN_SORT = [
    "AUTH", "DEV", "GLOB",
    "LOAN-DEC", "LOAN-DSB", "LOAN-RVW", "LOAN-GRD",
    "LOAN-INF", "LOAN-STAT", "LOAN-MGR", "LOAN-MBZD",
    "LOAN-APP", "LOAN-EXC", "LOAN-PRD", "LOAN-STP", "LOAN-ETC",
    "MBZ", "NOTIF", "RPT", "TERM", "USR", "UTIL",
]
LAYER_SORT = ["C", "S", "CV", "CL", "BT", "SCH", "FT", "EV", "UTL"]

# ── 분류 함수 ───────────────────────────────────────────────────

def get_module(fp):
    return "A" if "sofit-admin" in fp else "U"

def get_pkg_domain(fp):
    parts = fp.replace("\\", "/").split("/")
    try:
        idx  = parts.index("domain") + 1
        part = parts[idx] if idx < len(parts) else "global"
        return "global" if "." in part else part
    except ValueError:
        return "global"

def get_domain_code(cname, fp):
    pkg = get_pkg_domain(fp)
    if pkg == "loan":
        for prefix, code in LOAN_SUBDOMAIN:
            if cname.startswith(prefix):
                return f"LOAN-{code}"
        return "LOAN-ETC"
    if pkg == "global":
        # global 패키지지만 클래스명이 Loan 관련인 경우 (LoanDecisionProcessor 등)
        for prefix, code in LOAN_SUBDOMAIN:
            if cname.startswith(prefix):
                return f"LOAN-{code}"
        return "GLOB"
    return PKG_DOMAIN.get(pkg, "GLOB")

def get_layer(cname):
    for suffix, code in LAYER_MAP:
        if cname.endswith(suffix):
            return code
    return "UTL"

def sort_key(info):
    m = 0 if info["module"] == "A" else 1
    d = DOMAIN_SORT.index(info["domain"]) if info["domain"] in DOMAIN_SORT else 99
    l = LAYER_SORT.index(info["layer"])   if info["layer"]  in LAYER_SORT  else 99
    return (m, d, l, info["class_name"])

# ── 시나리오 / 조건 / 예상결과 추출 ──────────────────────────────

EXCEPTION_KWS = [
    "예외", "던진다", "실패", "없는", "없으면", "없을 때", "누락",
    "빈 문자", "blank", "Blank", "오류",
    "400", "403", "404", "409", "500",
    "권한이 없", "존재하지", "이미 ", "불가", "않으면", "않은",
    "throw", "Throw",
]

def is_exc(display):
    low = display.lower()
    return any(k.lower() in low for k in EXCEPTION_KWS)

def build_scenario(nested_disp, test_disp):
    tag = "[예외]" if is_exc(test_disp) else "[정상]"
    if nested_disp:
        return f"{tag} {nested_disp} - {test_disp}"
    return f"{tag} {test_disp}"

def extract_conditions(body):
    conds = []
    for m in re.finditer(r'given\(([^)]+)\)', body):
        obj = m.group(1).strip().split(".")[0][:40]
        if obj and obj not in conds:
            conds.append(obj)
    if not conds:
        m = re.search(r'\.\w+\((\d+L|"[^"]*"|[A-Z_]+\.\w+)\)', body)
        if m:
            conds.append(f"args: {m.group(1)}")
    return " / ".join(conds[:3]) or "-"

def extract_expected(body):
    parts = []
    sm = re.search(r'status\(\)\.(is\w+)\(\)', body)
    if sm:
        http = {
            "isOk": "200 OK", "isCreated": "201 Created",
            "isBadRequest": "400", "isUnauthorized": "401",
            "isForbidden": "403", "isNotFound": "404",
            "isConflict": "409", "isInternalServerError": "500",
        }
        parts.append(http.get(sm.group(1), sm.group(1)))

    cm = re.search(r'\.value\("([A-Z_][A-Z_0-9]+\d+)"\)', body)
    if cm:
        parts.append(f"code: {cm.group(1)}")

    em = re.search(r'isEqualTo\((\w+ErrorCode\.\w+)\)', body)
    if em:
        parts.append(f"예외: {em.group(1)}")

    am = re.search(r'assertThat\([^)]+\)\s*\.isEqualTo\(([^)]{1,50})\)', body)
    if am and "ErrorCode" not in am.group(1):
        parts.append(f"반환: {am.group(1).strip()}")

    vm = re.search(r'verify\([^)]+\)\.(\w+)\(', body)
    if vm:
        parts.append(f".{vm.group(1)}() 호출")

    return " / ".join(parts) or "-"

def get_importance(layer, domain, exc):
    if layer in ("C", "S"):
        if exc or any(d in domain for d in ["AUTH", "DEC", "MGR"]):
            return "높음"
        return "중간"
    if layer == "CV":
        return "낮음"
    if exc:
        return "중간"
    return "중간"

# ── Java 파서 ───────────────────────────────────────────────────

def strip_literals(text):
    """문자열 리터럴 제거 (줄 수 유지, 중괄호 오계산 방지)"""
    def _tb(m):
        return '""' + "\n" * m.group(0).count("\n")    # text block - 줄 수 보존
    text = re.sub(r'"""[\s\S]*?"""', _tb, text)
    text = re.sub(r'"(?:[^"\\]|\\.)*"', '""', text)    # regular string
    return text

def parse_java_file(fp):
    try:
        raw = open(fp, encoding="utf-8").read()
    except Exception as e:
        print(f"  [WARN] {Path(fp).name}: {e}")
        return []

    # 블록 주석 제거 → 문자열 리터럴 제거 → 라인 주석 제거 순서
    # ("http://..." 같은 문자열이 라인 주석 처리에 걸리는 것을 방지)
    no_block = re.sub(r'/\*[\s\S]*?\*/', '', raw)
    content  = re.sub(r'//[^\n]*', '', no_block)   # 표시용 (DisplayName 추출)
    clean    = strip_literals(no_block)             # 중괄호 계산용: 리터럴 먼저 제거
    clean    = re.sub(r'//[^\n]*', '', clean)       # 그 다음 라인 주석

    lines  = content.split("\n")
    clines = clean.split("\n")

    results  = []
    depth    = 0
    stack    = []   # {'depth', 'name', 'display', 'nested'}
    p_disp   = None
    p_nested = False
    p_test   = False

    i = 0
    while i < len(lines):
        raw_l = lines[i]
        cln_l = clines[i]
        s     = raw_l.strip()

        # @DisplayName
        dm = re.search(r'@DisplayName\("([^"]+)"\)', s)
        if dm:
            p_disp = dm.group(1)

        if "@Nested" in s:
            p_nested = True

        if re.search(r'@Test\b', s):
            p_test = True

        # 클래스 선언
        cm = re.search(r'\bclass\s+(\w+)', s)
        if cm and "@interface" not in s and "interface " not in s:
            entry = {
                "depth":   depth,
                "name":    cm.group(1),
                "display": p_disp or cm.group(1),
                "nested":  p_nested,
            }
            if not stack or p_nested:
                stack.append(entry)
            p_disp   = None
            p_nested = False
            p_test   = False  # 클래스 선언 후 @Test 플래그 리셋

        # 테스트 메서드
        elif p_test:
            vm = re.search(r'\bvoid\s+(\w+)\s*\(', s)
            if vm:
                method  = vm.group(1)
                display = p_disp or method
                p_disp  = None
                p_test  = False

                # 메서드 본문 수집
                body_parts  = []
                local_depth = 0
                j = i
                # 여는 중괄호 찾기
                while j < len(lines):
                    body_parts.append(lines[j].strip())
                    local_depth += clines[j].count("{") - clines[j].count("}")
                    j += 1
                    if local_depth > 0:
                        break
                # 닫는 중괄호까지 수집
                while j < len(lines) and local_depth > 0:
                    body_parts.append(lines[j].strip())
                    local_depth += clines[j].count("{") - clines[j].count("}")
                    j += 1

                body = "\n".join(body_parts)

                nested = [e for e in stack if e["nested"]]
                results.append({
                    "outer_class":    stack[0]["name"]     if stack  else "",
                    "nested_display": nested[-1]["display"] if nested else "",
                    "method":         method,
                    "display":        display,
                    "body":           body,
                    "filepath":       fp,
                })

        depth += cln_l.count("{") - cln_l.count("}")

        # 닫힌 스코프 팝
        while stack and depth <= stack[-1]["depth"]:
            stack.pop()

        i += 1

    return results

# ── 파일 수집 ───────────────────────────────────────────────────

def collect_files(test_dir):
    result = []
    for root, _, fnames in os.walk(test_dir):
        for fn in sorted(fnames):
            if not fn.endswith("Test.java"):
                continue
            if "Integration" in fn or "Real" in fn:
                continue
            result.append(os.path.join(root, fn))
    return result

# ── Excel 생성 ──────────────────────────────────────────────────

def mk_fill(hex_c):
    return PatternFill(start_color=hex_c, end_color=hex_c, fill_type="solid")

def mk_font(bold=False, color="000000", size=10):
    return Font(bold=bold, color=color, size=size, name="맑은 고딕")

def mk_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

HEADERS  = ["No", "기능 ID", "테스트 시나리오", "테스트 조건",
            "예상 결과", "단위 테스트 결과", "Fail 동작", "이슈 중요도"]
COL_WIDTHS = [5, 28, 55, 40, 35, 14, 20, 10]

def write_excel(rows):
    wb = openpyxl.Workbook()

    # ── 메인 시트 ──
    ws = wb.active
    ws.title = "단위테스트 명세서"

    # 헤더
    for c, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(1, c, h)
        cell.fill      = mk_fill(HEADER_BG)
        cell.font      = mk_font(True, "FFFFFF", 11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = mk_border()
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 28

    # 데이터
    for r, row in enumerate(rows, 2):
        color = LAYER_COLOR.get(row["layer"], "FFFFFF")
        vals  = [
            row["no"], row["id"], row["scenario"], row["conditions"],
            row["expected"], "PASS", "-", row["importance"],
        ]
        for c, v in enumerate(vals, 1):
            cell           = ws.cell(r, c, v)
            cell.border    = mk_border()
            cell.alignment = Alignment(vertical="center", wrap_text=True)

            if c == 6:  # PASS 셀
                cell.fill      = mk_fill(PASS_BG)
                cell.font      = mk_font(True, PASS_FG, 10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.fill = mk_fill(color)
                cell.font = mk_font(size=10)
                if c == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 45

    ws.freeze_panes = "A2"

    # ── 요약 시트 ──
    ws2 = wb.create_sheet("요약")
    sum_headers = ["모듈", "도메인", "레이어", "케이스 수", "PASS", "FAIL"]
    for c, h in enumerate(sum_headers, 1):
        cell           = ws2.cell(1, c, h)
        cell.fill      = mk_fill(HEADER_BG)
        cell.font      = mk_font(True, "FFFFFF", 10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = mk_border()
        ws2.column_dimensions[get_column_letter(c)].width = 18
    ws2.row_dimensions[1].height = 25

    counter = Counter((row["module"], row["domain"], row["layer"]) for row in rows)

    def summary_sort(item):
        (mod, dom, lay), _ = item
        m = 0 if mod == "A" else 1
        d = DOMAIN_SORT.index(dom) if dom in DOMAIN_SORT else 99
        l = LAYER_SORT.index(lay) if lay in LAYER_SORT else 99
        return (m, d, l)

    r2 = 2
    for (mod, dom, lay), cnt in sorted(counter.items(), key=summary_sort):
        ws2.cell(r2, 1, "Admin" if mod == "A" else "User")
        ws2.cell(r2, 2, dom)
        ws2.cell(r2, 3, lay)
        ws2.cell(r2, 4, cnt)
        ws2.cell(r2, 5, cnt)
        ws2.cell(r2, 6, 0)
        for c in range(1, 7):
            ws2.cell(r2, c).border    = mk_border()
            ws2.cell(r2, c).alignment = Alignment(horizontal="center", vertical="center")
            ws2.cell(r2, c).font      = mk_font(size=10)
        r2 += 1

    # 합계 행
    total = len(rows)
    for c, v in enumerate(["합계", "", "", total, total, 0], 1):
        cell           = ws2.cell(r2, c, v)
        cell.fill      = mk_fill(TOTAL_BG)
        cell.font      = mk_font(True, size=10)
        cell.border    = mk_border()
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws2.freeze_panes = "A2"

    wb.save(OUTPUT)
    print(f"\n저장 완료: {OUTPUT}")
    print(f"총 테스트 케이스: {total}개")

# ── Main ────────────────────────────────────────────────────────

def main():
    all_files = collect_files(ADMIN_TEST) + collect_files(USER_TEST)
    print(f"테스트 파일 {len(all_files)}개 발견\n")

    # 분류 후 정렬
    file_infos = []
    for fp in all_files:
        cname  = Path(fp).stem
        module = get_module(fp)
        domain = get_domain_code(cname, fp)
        layer  = get_layer(cname)
        file_infos.append({
            "fp": fp, "class_name": cname,
            "module": module, "domain": domain, "layer": layer,
        })
    file_infos.sort(key=sort_key)

    id_counter = defaultdict(int)
    rows = []
    no   = 1

    for info in file_infos:
        fp     = info["fp"]
        cname  = info["class_name"]
        module = info["module"]
        domain = info["domain"]
        layer  = info["layer"]

        test_cases = parse_java_file(fp)
        print(f"  [{module}] {domain}-{layer}  {cname}: {len(test_cases)}건")

        for tc in test_cases:
            key     = f"{module}-{domain}-{layer}"
            id_counter[key] += 1
            func_id = f"UT-{module}-{domain}-{layer}-{id_counter[key]:03d}"

            exc        = is_exc(tc["display"])
            scenario   = build_scenario(tc["nested_display"], tc["display"])
            conditions = extract_conditions(tc["body"])
            expected   = extract_expected(tc["body"])
            importance = get_importance(layer, domain, exc)

            rows.append({
                "no":         no,
                "id":         func_id,
                "module":     module,
                "domain":     domain,
                "layer":      layer,
                "scenario":   scenario,
                "conditions": conditions,
                "expected":   expected,
                "importance": importance,
            })
            no += 1

    print(f"\n총 {len(rows)}건 추출")
    write_excel(rows)

if __name__ == "__main__":
    main()
